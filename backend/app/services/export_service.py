"""导出服务（R16）：txt / docx，支持章节范围。"""

from __future__ import annotations

import io
import re
import sqlite3
import tempfile
import zipfile
from contextlib import closing
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from app.db.registry import BookRegistry, get_registry
from app.errors import ValidationError
from app.logging_config import get_logger, log_fields
from app.repositories import chapter_repo
from app.services import workspace
from app.utils.text import strip_html

logger = get_logger(__name__)

_RANGE_RE = re.compile(r"^\s*(\d+)\s*(?:-\s*(\d+)\s*)?$")

# 标题里已有的章号前缀，如「第1章」「第 12 章：」「第3章、」。
# 本项目章节 title 常被自动填成「第1章」这类题名，导出时必须先剥掉，
# 否则会渲染成「第1章 第1章」（P1）。数字一律以 row["seq"] 为准。
_CHAPTER_PREFIX_RE = re.compile(r"^\s*第\s*\d+\s*章\s*[：:、.．\-—]?\s*")


def parse_range(text: str | None) -> tuple[int | None, int | None]:
    if not text:
        return None, None
    m = _RANGE_RE.match(text)
    if not m:
        raise ValidationError(f"章节范围格式不正确：{text}（应为 1-10 或 5）")
    start = int(m.group(1))
    end = int(m.group(2)) if m.group(2) else start
    if start < 1 or end < 1 or end < start:
        raise ValidationError(f"章节范围不合法：{text}")
    return start, end


def _activate_book(slug: str) -> BookRegistry:
    """校验作品存在并把它设为当前活动库，返回 registry 供调用方取连接/目录。

    `require` 必须先于 `set_active`：后者会把 slug 落进共享的 active_book.json
    全局指针，若指向不存在的书就等于写坏这个指针。
    """
    registry = get_registry()
    registry.require(slug)
    workspace.set_active(slug)
    return registry


def _load_chapters(slug: str, range_text: str | None) -> tuple[list[dict], str]:
    registry = _activate_book(slug)
    start, end = parse_range(range_text)
    with registry.database(slug).connection() as conn:
        rows = chapter_repo.list_full_chapters(conn, start, end)
    meta = registry.read_meta(slug)
    return rows, (meta.get("title") or slug)


def _chapter_heading(row: dict) -> str:
    """章节标题：**只输出一个章号**，章号取自 `row["seq"]`。

    四种输入都收敛到同一形态（title 自带的章号一律丢弃，以 seq 为准）：
      ""            → 第5章
      "第1章"        → 第5章          （仅含章号）
      "第1章 风雪夜"  → 第5章 风雪夜   （剥掉旧前缀）
      "风雪夜"       → 第5章 风雪夜   （无章号，直接加前缀）
    """
    title = _CHAPTER_PREFIX_RE.sub("", (row.get("title") or "").strip()).strip()
    if not title:
        return f"第{row['seq']}章"
    return f"第{row['seq']}章 {title}"


def export_txt(slug: str, range_text: str | None) -> tuple[bytes, str]:
    rows, book_title = _load_chapters(slug, range_text)
    blocks: list[str] = [book_title, ""]
    for row in rows:
        blocks.append(_chapter_heading(row))
        blocks.append("")
        blocks.append(strip_html(row.get("content") or "").strip())
        blocks.append("")
    text = "\n".join(blocks).strip() + "\n"
    logger.info("export txt", **log_fields(slug=slug, chapters=len(rows)))
    return text.encode("utf-8"), f"{book_title}.txt"


def export_docx(slug: str, range_text: str | None) -> tuple[bytes, str]:
    from docx import Document

    rows, book_title = _load_chapters(slug, range_text)
    doc = Document()
    doc.add_heading(book_title, level=0)
    for row in rows:
        doc.add_heading(_chapter_heading(row), level=1)
        body = strip_html(row.get("content") or "").strip()
        for para in [p for p in body.split("\n") if p.strip()]:
            doc.add_paragraph(para)
    buffer = io.BytesIO()
    doc.save(buffer)
    logger.info("export docx", **log_fields(slug=slug, chapters=len(rows)))
    return buffer.getvalue(), f"{book_title}.docx"


def content_disposition(filename: str) -> str:
    ascii_fallback = re.sub(r"[^A-Za-z0-9_.-]", "_", filename) or "export"
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"


def export_backup(slug: str) -> tuple[bytes, str]:
    """整本 zip 备份（R16 扩展）：把"备份 = 复制整个书目录"的承诺产品化。

    库文件走 SQLite **在线备份 API** 取一致快照，而不是拷贝活文件 ——
    用户正在写作（本产品最高频场景）时，活 db + -wal 是不同瞬间的拼凑，
    换机恢复可能静默丢最近编辑甚至撕裂；快照则天然一致。
    打包只走白名单（novel.db 快照 + meta.json + exports/），
    避免 `meta.json.<uuid>.tmp` 原子写残留与崩溃伴生文件进包。
    """
    registry = _activate_book(slug)
    book_dir = registry.book_dir(slug)

    with (
        tempfile.TemporaryDirectory(prefix="ainovel-backup-") as tmp,
        registry.database(slug).connection() as conn,
    ):
        snap = Path(tmp) / "novel.db"
        with closing(sqlite3.connect(snap)) as dest:
            conn.backup(dest)  # 在线备份：自带读锁协调，含未合并的 WAL 内容

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(snap, "novel.db")
            meta_file = book_dir / "meta.json"
            if meta_file.is_file():
                zf.write(meta_file, "meta.json")
            exports_dir = book_dir / "exports"
            if exports_dir.is_dir():
                for item in sorted(exports_dir.rglob("*")):
                    if item.is_file():
                        zf.write(item, f"exports/{item.relative_to(exports_dir).as_posix()}")

    meta = registry.read_meta(slug)
    title = meta.get("title") or slug
    stamp = datetime.now().strftime("%Y%m%d")
    logger.info("export backup", **log_fields(slug=slug, bytes=buf.tell()))
    return buf.getvalue(), f"{title}-备份-{stamp}.zip"
